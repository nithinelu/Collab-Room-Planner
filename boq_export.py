"""Excel export for Collab Room Planner.

One BOQ sheet per option (Item / Description / Qty columns, so the BOQ Costing Tool can read it
as a plain hardware BOQ), then a Summary sheet with room data, coverage and the 3D snapshots.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD_FILL = PatternFill("solid", fgColor="4F81BD")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _head(ws, row, values):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(vertical="center")


def write_workbook(path, payload, decode):
    room = payload["room"]
    name = payload.get("name") or "Room"
    wb = Workbook()
    wb.remove(wb.active)

    for opt in payload["options"]:
        title = f"BOQ {opt['tag']} - {opt['vendor']}"[:31]
        ws = wb.create_sheet(title)
        _head(ws, 1, ["S/N", "Item", "Description", "Qty"])
        for i, (item, qty, note) in enumerate(opt["rows"], 1):
            for col, v in enumerate([i, item, note, qty], 1):
                c = ws.cell(row=i + 1, column=col, value=v)
                c.border = BOX
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 62
        ws.column_dimensions["C"].width = 44
        ws.column_dimensions["D"].width = 8
        ws.freeze_panes = "A2"

    for opt in payload["options"]:
        if not opt.get("runs"):
            continue
        ws = wb.create_sheet(f"Install {opt['tag']} - {opt['vendor']}"[:31])
        ws["A1"] = f"{name} · {opt['vendor']} {opt['device']} · installation"
        ws["A1"].font = Font(bold=True, size=13)
        cab = opt.get("cabling") or {}
        ws["A2"] = (f"Route: {cab.get('route', '')} · Equipment: {cab.get('rack', '')} · "
                    f"+{cab.get('slack', 0)}% slack, {cab.get('loop', 0)} m service loop each end")
        ws["A2"].font = Font(italic=True, color="666666")
        r = 4
        _head(ws, r, ["Mounting", "Value", "Check", "Note"])
        for item, value, check, note in opt["mount"]:
            r += 1
            for col, v in enumerate([item, value, {"ok": "OK", "warn": "Check", "bad": "Fail"}.get(check, ""), note], 1):
                ws.cell(row=r, column=col, value=v).border = BOX
            if check in ("warn", "bad"):
                ws.cell(row=r, column=3).font = Font(bold=True, color="C0392B" if check == "bad" else "B7791F")
        r += 2
        _head(ws, r, ["#", "From", "To", "Cable", "Purpose", "Length (m)"])
        first = r + 1
        for row in opt["runs"]:
            r += 1
            for col, v in enumerate(row, 1):
                ws.cell(row=r, column=col, value=v).border = BOX
        r += 1
        ws.cell(row=r, column=5, value="Total (all cables)").font = Font(bold=True)
        ws.cell(row=r, column=6, value=f"=SUM(F{first}:F{r - 1})").font = Font(bold=True)
        r += 1
        ws.cell(row=r, column=5, value="CAT6 total").font = Font(bold=True)
        ws.cell(row=r, column=6, value=f'=SUMIF(D{first}:D{r - 2},"CAT6",F{first}:F{r - 2})').font = Font(bold=True)
        pw = opt.get("power")
        if pw:
            r += 2
            _head(ws, r, ["PoE peripheral", "PoE standard", "Class W", "Typical W", "Powered from"])
            for row in pw["rows"]:
                r += 1
                for col, v in enumerate(row, 1):
                    ws.cell(row=r, column=col, value=v).border = BOX
            for line in pw["summary"]:
                r += 1
                ws.cell(row=r, column=1, value=line).font = Font(italic=True)
        r += 2
        ws.cell(row=r, column=1, value="Power points: " + cab.get("power", ""))
        for col, w in zip("ABCDEF", (34, 30, 22, 14, 40, 12)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A5"

    ws = wb.create_sheet("Summary")
    ws["A1"] = name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Collab Room Planner · {date.today():%d %b %Y}"
    ws["A2"].font = Font(italic=True, color="666666")
    rows = [
        ("Room type", room.get("layout", "Boardroom")),
        ("Room width (m)", room["W"]), ("Room depth from display (m)", room["D"]),
        ("Ceiling height (m)", room["H"]), ("Seats", room["seats"]), ("Table width (m)", room["tw"]),
        ("Display to table (m)", room["gap"]), ("Mic planning radius (%)", room["plan"]),
        ("Dual displays", "Yes" if room.get("dual") else "No"),
    ]
    _head(ws, 4, ["Room", "Value"])
    for i, (k, v) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=k).border = BOX
        ws.cell(row=i, column=2, value=v).border = BOX

    r = 5 + len(rows) + 1
    opts = payload["options"]
    _head(ws, r, ["Coverage"] + [f"{o['tag']} · {o['vendor']} {o['device']}" for o in opts])
    metrics = [
        ("Seats fully covered", lambda s: f"{s['ok']}/{s['total']}"),
        ("Marginal seats", lambda s: s["warn"]),
        ("Seats not covered", lambda s: s["bad"]),
        ("Camera coverage", lambda s: f"{s['camera']}/{s['total']}"),
        ("Audio coverage", lambda s: f"{s['audio']}/{s['total']}"),
        ("Face-on across table", lambda s: f"{s.get('face', '-')}/{s['total']}"),
        ("Extra cameras", lambda s: s.get("extra_cams", "None")),
        ("Minimum display (in)", lambda s: s["display"]),
        ("HFOV (°)", lambda s: s["hfov"]),
        ("Camera reach (m)", lambda s: s["reach"]),
        ("Mic pickup (m)", lambda s: s["mic"]),
        ("Add-on mics", lambda s: s["addon_mics"]),
    ]
    for i, (label, fn) in enumerate(metrics, r + 1):
        ws.cell(row=i, column=1, value=label).border = BOX
        for j, o in enumerate(opts, 2):
            ws.cell(row=i, column=j, value=fn(o["summary"])).border = BOX
    r += len(metrics) + 2
    for o in opts:
        ws.cell(row=r, column=1, value=f"Notes {o['tag']}").font = Font(bold=True)
        for note in o["notes"]:
            r += 1
            ws.cell(row=r, column=1, value="• " + note)
        r += 2

    ws.column_dimensions["A"].width = 34
    for col in "BC":
        ws.column_dimensions[col].width = 36

    # 3D snapshots, stacked under the notes
    for o in opts:
        if not o.get("image"):
            continue
        ws.cell(row=r, column=1, value=f"3D view {o['tag']} · {o['vendor']} {o['device']}").font = Font(bold=True)
        img = XLImage(io.BytesIO(decode(o["image"])))
        img.width, img.height = 768, 432
        ws.add_image(img, f"A{r + 1}")
        r += 24

    wb.save(path)
